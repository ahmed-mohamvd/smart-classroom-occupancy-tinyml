from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
HEADER_IMAGE = ROOT / "utm_fee_header.png"


def set_run_font(run, size=10.3, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, data in edges.items():
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        for key, value in data.items():
            node.set(qn("w:" + key), str(value))


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_header_block(doc, week_number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(HEADER_IMAGE), width=Inches(4.25))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("VECAD ELITE INTERNSHIP 2025"), 12, bold=False)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(13)
    set_run_font(p.add_run(f"WEEK {week_number}"), 12, bold=False)

    details = doc.add_table(rows=3, cols=2)
    details.alignment = WD_TABLE_ALIGNMENT.LEFT
    details.autofit = False
    details.columns[0].width = Inches(1.55)
    details.columns[1].width = Inches(4.95)
    values = [
        ("NAME", "Ahmed Mohamed Abdalla Ahmed"),
        ("MATRIC", "A24KE0429"),
        ("SUPERVISOR", "Dr Mohd Anuar Md Ali"),
    ]
    for row, (label, value) in zip(details.rows, values):
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for cell in row.cells:
            set_cell_border(cell, top={"val": "nil"}, bottom={"val": "nil"}, left={"val": "nil"}, right={"val": "nil"})
            set_cell_margin(cell, top=0, bottom=0, start=0, end=0)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
        set_run_font(row.cells[0].paragraphs[0].add_run(label), 11.5)
        set_run_font(row.cells[1].paragraphs[0].add_run(": " + value), 11.5)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), 11.5, bold=True)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.28)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.04
    p.paragraph_format.widow_control = True
    set_run_font(p.add_run(text), 10.3)


def add_signature_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.25)
    table.columns[1].width = Inches(3.25)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margin(cell, top=115, start=120, bottom=110, end=120)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "9", "color": "000000"},
            bottom={"val": "single", "sz": "9", "color": "000000"},
            left={"val": "single", "sz": "9", "color": "000000"},
            right={"val": "single", "sz": "9", "color": "000000"},
        )
    labels = [
        ("Student's Name/Signature: ", "____________________________", "Date: ", "__________________"),
        ("Supervisor's Name/Signature: ", "________________________", "Date: ", "__________________"),
    ]
    for cell, (a, b, c, d) in zip(table.rows[0].cells, labels):
        p1 = cell.paragraphs[0]
        p1.paragraph_format.space_after = Pt(20)
        set_run_font(p1.add_run(a), 9.6, bold=True)
        set_run_font(p1.add_run(b), 9.6)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(c), 9.6, bold=True)
        set_run_font(p2.add_run(d), 9.6)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.30)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(10.3)
    style.font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.author = "Ahmed Mohamed Abdalla Ahmed"


def build_logbook(week_number, objective_paragraphs, page_two_paragraphs, page_three_paragraphs, output_name, subject):
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = f"VECAD ELITE Internship Week {week_number} Logbook"
    doc.core_properties.subject = subject

    add_header_block(doc, week_number)
    add_section_heading(doc, "Objectives:")
    for paragraph in objective_paragraphs:
        add_body(doc, paragraph)
    add_signature_block(doc)

    doc.add_page_break()
    add_header_block(doc, week_number)
    add_section_heading(doc, "Summary of weekly activities:")
    for paragraph in page_two_paragraphs:
        add_body(doc, paragraph)
    add_signature_block(doc)

    doc.add_page_break()
    add_header_block(doc, week_number)
    add_section_heading(doc, "Summary of weekly activities:")
    for paragraph in page_three_paragraphs:
        add_body(doc, paragraph)
    add_signature_block(doc)

    output = ROOT / output_name
    doc.save(output)
    print(output)


WEEK_2_OBJECTIVES = [
    "The main objective for the second internship week was to convert the first-week design into a working multi-sensor prototype. I focused on assembling the available modules, checking the quality of soldered connections and wiring, and verifying every sensor through controlled tests before relying on the full system. The goal was not simply to make a display show values; it was to prove that each reading was obtained from the intended component and could be communicated reliably to the STM32 NUCLEO-L476RG board.",
    "A second objective was to establish a dependable I2C hardware and software layer. The BME280, BH1750, VL53L0X, and SSD1306 OLED were intended to share the same I2C bus, so each device needed to respond at its expected address without disturbing the others. I aimed to use an I2C scanner, module-specific test programs, and serial output to distinguish between software issues, incorrect connections, insufficient power, pull-up problems, and an actual faulty module. This process was necessary before integration because a single defective or incorrectly supplied module can prevent the shared bus from operating correctly.",
    "The final objective was to prepare the prototype for structured occupancy data collection. This required adding the two PIR sensors, LEDs, push buttons, and active buzzer; defining the EMPTY, LOW, and HIGH occupancy labels; and confirming that the firmware could start and stop labelled recording sessions. I also aimed to prepare the serial logging method so that measurements could be saved as CSV files for later data cleaning and AI training, rather than being lost in the serial monitor."
]

WEEK_2_PAGE_2 = [
    "At the beginning of the week, the ordered components became available and I started the practical assembly stage. The work took more time than expected because I did not initially have all of the required soldering tools and had to prepare or borrow the appropriate equipment before the headers could be attached properly. I treated this stage carefully because poor solder joints can create intermittent faults that look like software problems. After fitting the headers, I inspected the modules, connected them one at a time, and avoided integrating all of them at once. This made it possible to identify the source of a problem instead of having many unknowns in the same circuit.",
    "The BME280 environmental sensor was one of the first modules that I tested. I confirmed that it could return temperature, relative humidity, and pressure values and used the serial output to check that the data changed in a reasonable way. The OLED display was then brought up using the SSD1306 driver, and I confirmed that measured values and simple status messages could be displayed. These first tests established that the STM32 I2C peripheral, the basic wiring, and the OLED library could work together. They also gave me a useful way to show live system status without depending only on a computer terminal.",
    "I developed and used an I2C scanner to examine every valid address on the selected STM32 I2C bus. The scanner was especially useful because it reported the normal seven-bit device address as well as the left-shifted address required by the STM32 HAL functions. When modules were connected correctly, the expected devices could be recognised from their addresses: the BH1750 at 0x23, the VL53L0X at 0x29, the SSD1306 OLED at 0x3C, and the BME280 at 0x76. This gave a clear and repeatable hardware check before running each sensor's application code.",
    "The VL53L0X distance sensor required additional firmware work because its vendor API includes platform files intended for several operating systems. I selected the STM32-compatible I2C platform implementation and removed the Windows-specific communication and logging sources that caused compilation errors. After the driver was built successfully, the sensor reported distance and range status values through the serial monitor and the OLED display. I learned to interpret a valid range status separately from an invalid or out-of-range reading, rather than treating every numeric value as a real distance."
]

WEEK_2_PAGE_3 = [
    "The BH1750 light-sensor test became an important troubleshooting exercise. The first module did not answer on the I2C bus even after the wiring, pull-up configuration, supply, and addresses had been checked on both the NUCLEO board and an Arduino. Instead of continuing to modify code without evidence, I compared it with another known working sensor and repeated the scan. The replacement BH1750 was detected at address 0x23 and produced sensible lux values that changed when the lighting level changed. This result showed that the original issue was most likely the module itself rather than the I2C scanner or the STM32 software. It also reinforced the value of isolating a suspected device and testing it with a second platform.",
    "I also tested the two PIR motion sensors, the LEDs, push buttons, and active buzzer. The PIR modules provided a digital logic signal when motion was detected, but their behaviour had to be observed after the warm-up period and under real room conditions. I learned that their delay, sensitivity, mounting direction, and air movement can affect the output, so a single PIR transition should not automatically be treated as a reliable occupancy decision. The active buzzer was verified with a simple GPIO test program. This test confirmed the logic level required by the particular module and prevented the main program from leaving the buzzer on continuously because of an inverted active-low connection.",
    "After the individual checks, I combined the sensor modules into one prototype and confirmed that the shared I2C bus could identify all four required devices. The main firmware was extended to read the two PIR inputs, BME280 values, ambient light, and distance information while also driving the OLED, LEDs, buttons, and buzzer. I introduced the three recording labels - EMPTY, LOW, and HIGH - so that a session could be given its correct class before logging began. The LEDs and display made the selected state visible, while the buttons allowed the operator to change the label and start or stop a recording session without repeatedly editing the code.",
    "To support later AI work, I prepared a serial-data logger in Python. The STM32 firmware emitted a CSV-style header and labelled data rows containing the session identifier, microcontroller time, PIR values, distance data, environmental readings, light level, validity flag, and occupancy label. The Python logger listened on the virtual COM port and saved each complete session as a CSV file. I then planned a balanced collection routine for EMPTY, LOW, and HIGH occupancy under different conditions such as lights on or off and air-conditioning on or off. By the end of the week, the prototype had moved from individual module tests to a stable integrated data-collection system."
]

WEEK_3_OBJECTIVES = [
    "The main objective for the third internship week was to transform the labelled sensor recordings into a small, testable AI occupancy classifier. This involved reviewing the collected CSV sessions, removing unsuitable transition data, creating meaningful input features from the PIR signals, and training a model that could distinguish EMPTY, LOW, and HIGH occupancy. The objective was not only to obtain a high accuracy value on a computer, but to build a model compact enough to be converted and executed on the STM32 microcontroller.",
    "A second objective was to understand and apply the complete TensorFlow and STM32Cube.AI deployment workflow. I aimed to document why the data needed normalisation, why sensor readings were grouped into time windows, how the neural-network layers produced three class probabilities, and how the trained model could be converted to an INT8 TensorFlow Lite file. This was essential for explaining the project in the presentation and for ensuring that the computer-side model and the embedded implementation used the same feature calculation and scaling.",
    "The final objective was to improve the prototype from a classifier demonstration into an energy-aware monitoring system. I planned to integrate the AI output with the OLED, LEDs, and buzzer and to use environmental readings for rule-based energy-saving warnings. At the same time, I reviewed the real limitation of the VL53L0X sensor in the cardboard prototype and documented why it should not be treated as a primary occupancy feature in the final model without further hardware improvement."
]

WEEK_3_PAGE_2 = [
    "The week began with a review of the completed prototype and the practical sensing limitations found during data collection. The VL53L0X sensor can provide a useful distance reading for an object located directly in front of it, but its field of view is narrow and its reliable working distance in the assembled prototype was about 80 cm. A person standing outside the direct line of sight, or remaining still at a greater distance, was therefore not detected consistently. Following the supervisor discussion, I kept the sensor as a monitored part of the system but did not make it a core input to the trained occupancy classifier. This was an engineering decision based on the observed evidence rather than a decision to force an unreliable feature into the AI model.",
    "I consolidated the labelled CSV recordings collected from EMPTY, LOW, and HIGH occupancy sessions. Before training, I reviewed the data so that moments when the operator approached the device to press the start or stop button would not be treated as ordinary room behaviour. The early and final transition rows of each session were removed where necessary, and invalid or incomplete rows were excluded. This cleaning step was important because an AI model learns from whatever examples it is given. If these transitions remained labelled as EMPTY, for example, the model could learn misleading patterns that would reduce its reliability during real operation.",
    "The cleaned measurements were converted into overlapping time windows. Each window contained 30 PIR samples, representing approximately 30 seconds of recent activity, and a new inference window was produced every five samples, or approximately every five seconds. From the two PIR sequences, I calculated six summary features that describe motion activity within the window. The features gave the model information about the amount and distribution of detected movement rather than asking it to interpret a single instantaneous PIR value. This approach is better suited to PIR sensing because occupancy is a pattern over time, while an individual PIR reading is only a momentary motion event.",
    "I normalised the features using the same limits derived from the training data. Normalisation placed features with different original scales into a comparable range before the neural network processed them. Without this step, a feature with a naturally larger numerical range could dominate the training process even if it was not more meaningful. The same normalisation values were later placed in the STM32 preprocessing code so that live features on the board were transformed in exactly the same way as the features used during TensorFlow training. This matching of the training and embedded preprocessing stages is essential for a valid deployment."
]

WEEK_3_PAGE_3 = [
    "Using TensorFlow, I trained a compact neural network with six input features, two hidden layers containing 16 and 8 ReLU neurons, and a three-output Softmax layer for EMPTY, LOW, and HIGH. The final network contained 275 parameters, which is small enough for an embedded target. The recorded windows were separated into training, validation, and test sets so that the final reported performance was measured on data that had not been used to update the model weights. I set a maximum of 300 epochs, but early stopping ended training at 49 epochs because the validation result was no longer improving. This prevented unnecessary training and helped reduce the risk of overfitting.",
    "The final floating-point model achieved a test accuracy of 88.67 percent. It was then converted to an INT8 TensorFlow Lite model, which occupied 3,616 bytes while preserving the same measured test accuracy. I analysed the INT8 model in STM32Cube.AI and generated the C source files, network data, weights, and configuration files required by the STM32CubeIDE project. This proved that the model was compatible with the STM32L476RG target and that the AI component was not only a desktop result. The generated model was then wrapped with application functions for initialisation and inference.",
    "In the STM32 firmware, I implemented a feature buffer that stores the latest 30 PIR samples. Once this first window is full, the system calculates the same six features used during training, normalises and quantises them, runs the AI network, and prints the three class probabilities through the serial monitor. Because the window advances by five samples, the system produces a new occupancy decision every five seconds after the initial window is available. I tested the live output with the existing push buttons, LEDs, OLED, and PIR wiring kept in place. The serial output showed the predicted class, confidence value, and probabilities for EMPTY, LOW, and HIGH, allowing the decision to be checked during demonstrations.",
    "Finally, I extended the prototype with energy-aware warning behaviour. The BH1750 and BME280 remained useful even though they were not used as features in the PIR-only occupancy model. Their readings were used by simple rule-based checks together with the AI EMPTY output; for example, the system can warn the user when the room appears empty while the light level remains high for a sustained period. The OLED, LEDs, and buzzer provide immediate feedback for these conditions. I also improved the physical arrangement by placing the two PIR sensors in opposite directions on the cardboard prototype to increase coverage. The remaining improvement path is clear: collect more varied quiet LOW and HIGH sessions, improve the physical sensor placement, and consider an additional PIR or a longer-range presence sensor before expanding the final AI feature set."
]


if __name__ == "__main__":
    build_logbook(
        2,
        WEEK_2_OBJECTIVES,
        WEEK_2_PAGE_2,
        WEEK_2_PAGE_3,
        "Week_2_E_Logbook_Ahmed.docx",
        "Hardware integration, sensor validation, and data collection preparation",
    )
    build_logbook(
        3,
        WEEK_3_OBJECTIVES,
        WEEK_3_PAGE_2,
        WEEK_3_PAGE_3,
        "Week_3_E_Logbook_Ahmed.docx",
        "AI training, STM32 deployment, and energy-aware system improvement",
    )
