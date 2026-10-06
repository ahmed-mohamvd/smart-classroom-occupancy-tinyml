from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Week_1_E_Logbook_Ahmed.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs[edge]
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in edge_data.items():
            element.set(qn("w:{}".format(key)), str(value))


def set_cell_margin(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=10.5, bold=False, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_body(paragraph, text):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.28)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.widow_control = True
    run = paragraph.add_run(text)
    set_run_font(run, 10.5)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, 12.5 if level == 1 else 11, bold=True)
    return paragraph


def add_title(doc, week_label, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("UTM FACULTY OF ELECTRICAL ENGINEERING")
    set_run_font(r, 10.5, bold=True)

    p = doc.add_paragraph()
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("VECAD ELITE INTERNSHIP 2025")
    set_run_font(r, 18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(week_label)
    set_run_font(r, 13, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(subtitle)
    set_run_font(r, 10.5, italic=True)


def add_details_table(doc):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.28)
    table.columns[1].width = Inches(3.28)
    data = [
        ("NAME", "Ahmed Mohamed Abdalla Ahmed"),
        ("MATRIC", "A24KE0429"),
        ("SUPERVISOR", ""),
        ("PROJECT", "Smart Classroom Occupancy and Energy Monitoring System"),
    ]
    idx = 0
    for row in table.rows:
        for cell in row.cells:
            label, value = data[idx]
            idx += 1
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "F2F2F2" if label != "PROJECT" else "FFFFFF")
            set_cell_margin(cell)
            set_cell_border(
                cell,
                top={"val": "single", "sz": "6", "color": "D9D9D9"},
                bottom={"val": "single", "sz": "6", "color": "D9D9D9"},
                left={"val": "single", "sz": "6", "color": "D9D9D9"},
                right={"val": "single", "sz": "6", "color": "D9D9D9"},
            )
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(label + ": ")
            set_run_font(r, 9.5, bold=True)
            r = p.add_run(value if value else "________________________________")
            set_run_font(r, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_signatures(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            set_cell_margin(cell, top=90, bottom=90)
            set_cell_border(
                cell,
                top={"val": "nil"},
                bottom={"val": "nil"},
                left={"val": "nil"},
                right={"val": "nil"},
            )
    labels = [
        "Student's Name/Signature: ____________________________",
        "Supervisor's Name/Signature: _________________________",
        "Date: ____________________",
        "Date: ____________________",
    ]
    k = 0
    for row in table.rows:
        for cell in row.cells:
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(labels[k])
            set_run_font(r, 9.2)
            k += 1


def add_continuation_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("VECAD ELITE INTERNSHIP 2025")
    set_run_font(r, 11, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("WEEK 1 - SUMMARY OF WEEKLY ACTIVITIES")
    set_run_font(r, 12.5, bold=True)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    add_title(doc, "WEEK 1", "Project Familiarization and Initial System Planning")
    add_details_table(doc)
    add_heading(doc, "Objectives")

    objectives = [
        "The main objective for my first internship week was to establish a clear and practical foundation for the Smart Classroom Occupancy and Energy Monitoring System before moving into hardware assembly or AI deployment. I needed to understand the capability of the STM32 NUCLEO-L476RG board, identify the sensing elements required by the proposed system, and convert the initial project idea into a structured engineering workflow. The system is intended to distinguish between an empty room, low occupancy, and high occupancy, while also monitoring environmental conditions that can support energy-saving decisions.",
        "A second objective was to develop a sensible hardware architecture. I planned how the PIR motion sensors, BME280 temperature and humidity sensor, BH1750 ambient-light sensor, VL53L0X distance sensor, SSD1306 OLED display, LEDs, push buttons, and active buzzer could be used together with the NUCLEO board. Particular attention was given to the I2C bus because several of the selected modules communicate through this protocol. I also aimed to prepare a basic testing strategy so that each component could be verified independently before all modules were integrated into one system.",
        "The final objective was to begin building the knowledge needed for the artificial intelligence phase of the project. I explored the relationship between sensor data collection, labelled occupancy classes, model training, TensorFlow, and STM32Cube.AI deployment. This early preparation was important because the project would eventually require not only working sensors but also a reliable path for converting recorded measurements into a compact model that can run directly on the STM32 microcontroller."
    ]
    for text in objectives:
        format_body(doc.add_paragraph(), text)
    add_signatures(doc)

    doc.add_page_break()
    add_continuation_title(doc)
    add_heading(doc, "Summary of Weekly Activities")

    page_two = [
        "During the first week, I received the STM32 NUCLEO-L476RG development board and began by confirming that the board was functioning correctly. I configured the development environment, connected the board through the ST-LINK interface, and carried out a simple LED blinking test. This was an important first check because it verified the basic programming and debugging connection before additional hardware was connected. After the initial test, I studied the main characteristics of the STM32L476RG microcontroller used on the board. I reviewed its operating speed, memory resources, GPIO capability, peripheral interfaces, and the layout of the Arduino and Morpho headers. This helped me understand which pins could be reserved for I2C communication, digital PIR inputs, LEDs, push buttons, and the buzzer output.",
        "I then translated the project idea into a preliminary operating plan. Rather than attempting to connect every device and train an AI model immediately, I organised the work into a sequence of engineering stages. The first stage was individual component testing, followed by full hardware integration, structured data collection, data preparation and model training, and finally deployment of the trained model to the STM32 target. This workflow provided a clear reason for every activity carried out later in the project. It also helped identify dependencies early: for example, reliable I2C communication and stable sensor readings were necessary before meaningful data could be collected, while balanced labelled data would be necessary before the AI model could be evaluated fairly.",
        "A preliminary circuit arrangement was designed for the selected components. The BME280, BH1750, VL53L0X, and OLED display were planned to share the I2C bus, while the PIR modules would use separate digital input pins because their output represents motion events. The LEDs and active buzzer were planned as user-feedback outputs, and two push buttons were reserved for basic interaction during recording and testing. I also considered how the system could be physically arranged so that the sensors would have a useful viewing direction and the wiring could remain manageable during experiments. The circuit design at this stage was intentionally treated as a working design that could be improved after real hardware testing revealed practical limitations.",
        "At the same time, I prepared the component procurement list. For each module I checked its intended purpose, operating voltage, interface type, and expected role in the occupancy-monitoring system. I researched how the PIR sensors detect movement, how the BME280 measures temperature, humidity, and pressure, how the BH1750 reports ambient light level in lux, and how the VL53L0X performs time-of-flight distance measurement. I also studied the SSD1306 OLED display because it would provide immediate feedback without requiring a computer during demonstrations. Reviewing these modules before their arrival allowed me to identify likely libraries, sample codes, and common wiring requirements in advance."
    ]
    for text in page_two:
        format_body(doc.add_paragraph(), text)

    doc.add_page_break()
    add_continuation_title(doc)
    add_heading(doc, "Summary of Weekly Activities Continued")

    page_three = [
        "A major learning activity in this week was the I2C communication protocol. Since several modules in the proposed system use I2C, I studied the roles of SDA and SCL, pull-up resistors, seven-bit device addresses, acknowledge signals, and the difference between a normal seven-bit address and the left-shifted address commonly used by STM32 HAL functions. I also prepared an I2C scanner approach that could later be used to confirm whether each connected device was responding on the shared bus. This preparation was useful because it provided a systematic way to diagnose future issues: before blaming a driver or an application program, I could first check power, ground, bus wiring, and the detected address of the component.",
        "While waiting for the components and preparing the first test codes, I began exploring the AI workflow required for an embedded system. I reviewed the purpose of TensorFlow as a software framework for creating and training machine-learning models and examined how STM32Cube.AI can convert a trained model into C code that can be executed on the STM32 microcontroller. To make this new topic more concrete, I studied and attempted a small reference project based on handwritten digit classification on an MCU. The reference was used as a learning exercise rather than as a direct solution for my project. It helped me understand the overall sequence of preparing data, defining a neural-network structure, training the model on a computer, converting it to TensorFlow Lite, analysing it in STM32Cube.AI, and integrating the generated files into an STM32CubeIDE project.",
        "This early AI research also made the importance of data quality clear. The future occupancy model would not learn directly from an idea such as 'a room is busy'; it would learn from labelled examples collected under different real conditions. Therefore, I began considering the three output labels that would later be used: EMPTY, LOW, and HIGH occupancy. I also noted that data would need to include environmental variation, such as changes in lighting and air-conditioning conditions, so that the model would not make decisions based on a single fixed room condition. This planning influenced the later data-collection design and helped connect the hardware work with the AI objective of the project.",
        "The week ended with the project plan, initial circuit design, component list, I2C test strategy, and first AI learning path ready for the next stage. The main challenge was that the project combined several areas that were new to me at the same time: STM32 firmware development, multi-sensor interfacing, practical circuit assembly, serial data handling, and embedded AI deployment. Instead of treating this as a single large task, I divided it into manageable tests and recorded what each stage needed to prove. This gave me a more realistic and organised starting point for the following week, when the components could be assembled, tested individually, and gradually integrated into the working prototype."
    ]
    for text in page_three:
        format_body(doc.add_paragraph(), text)
    add_signatures(doc)

    # Simple centered page numbering in the footer.
    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run("Week 1 | Page ")
        set_run_font(r, 8.5)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)

    doc.core_properties.title = "VECAD ELITE Internship Week 1 Logbook"
    doc.core_properties.author = "Ahmed Mohamed Abdalla Ahmed"
    doc.core_properties.subject = "Week 1 project familiarization and planning"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)
